import io
import os
import tempfile
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook
from reportlab.lib.utils import open_for_read

from app.core.security import hash_password
from app.models.audit_log import AuditLog
from app.models.enums import AuditAction, AuditEntityType, DecisionStatus, UserRole
from app.models.user import User


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _create_user(db_session, email, employee_id, role=UserRole.EMPLOYEE, department=None):
    user = User(
        full_name=f"Export User {email.split('@')[0]}",
        email=email,
        role=role,
        password=hash_password("password123"),
        employee_id=employee_id,
        department=department,
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


def _auth_headers(user, make_token):
    return {"Authorization": f"Bearer {make_token(str(user.id))}"}


def _create_decision(client, headers, title="Test Decision", category="Engineering"):
    return client.post(
        "/decisions",
        json={"title": title, "problem_statement": "PS", "category": category},
        headers=headers,
    )


def _create_audit_log(db_session, user_id, action, entity_type, entity_id=None, description=None, ip_address=None):
    entry = AuditLog(
        user_id=user_id,
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        description=description,
        ip_address=ip_address,
    )
    db_session.add(entry)
    db_session.commit()
    db_session.refresh(entry)
    return entry


# ===========================================================================
# PDF Export Tests
# ===========================================================================

class TestExportDecisionsPDF:
    def test_returns_pdf(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf_rpt@example.com", "DEMP_PDF", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_decision(client, headers, title="Decision A")

        res = client.get("/reports/decisions/export/pdf", headers=headers)
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"
        assert "decisions_report.pdf" in res.headers.get("content-disposition", "")
        assert len(res.content) > 0

    def test_pdf_starts_with_pdf_header(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf2_rpt@example.com", "DEMP_PDF2", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_decision(client, headers)

        res = client.get("/reports/decisions/export/pdf", headers=headers)
        assert res.content[:5] == b"%PDF-"

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "dempdf_rpt@example.com", "DEMP_EPDF", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)

        res = client.get("/reports/decisions/export/pdf", headers=headers)
        assert res.status_code == 200
        assert len(res.content) > 0

    def test_filter_by_status(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf3_rpt@example.com", "DEMP_PDF3", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers, title="Approved Dec").json()
        client.patch(f"/decisions/{r['id']}/status", json={"status": "Approved"}, headers=headers)
        _create_decision(client, headers, title="Draft Dec")

        res = client.get("/reports/decisions/export/pdf?status=Approved", headers=headers)
        assert res.status_code == 200
        assert len(res.content) > 0

    def test_filter_by_category(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf4_rpt@example.com", "DEMP_PDF4", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_decision(client, headers, category="Engineering")
        _create_decision(client, headers, category="Marketing")

        res = client.get("/reports/decisions/export/pdf?category=Engineering", headers=headers)
        assert res.status_code == 200

    def test_filter_by_date_range(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf5_rpt@example.com", "DEMP_PDF5", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_decision(client, headers)

        today = datetime.now().strftime("%Y-%m-%d")
        yesterday = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
        res = client.get(f"/reports/decisions/export/pdf?start_date={yesterday}&end_date={today}", headers=headers)
        assert res.status_code == 200

    def test_invalid_date_returns_422(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf6_rpt@example.com", "DEMP_PDF6", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/decisions/export/pdf?start_date=bad", headers=headers)
        assert res.status_code == 422

    def test_invalid_date_range_returns_422(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf7_rpt@example.com", "DEMP_PDF7", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/decisions/export/pdf?start_date=2026-12-31&end_date=2026-01-01", headers=headers)
        assert res.status_code == 422

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/decisions/export/pdf")
        assert res.status_code == 401

    def test_content_disposition_header(self, client, db_session, make_token):
        user = _create_user(db_session, "dpdf8_rpt@example.com", "DEMP_PDF8", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/decisions/export/pdf", headers=headers)
        assert "attachment" in res.headers["content-disposition"]


class TestExportApprovalsPDF:
    def test_returns_pdf(self, client, db_session, make_token):
        user = _create_user(db_session, "apdf_rpt@example.com", "AEMP_PDF", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers).json()
        _create_audit_log(db_session, user.id, "approve", "decision", r["id"], "Approved")

        res = client.get("/reports/approvals/export/pdf", headers=headers)
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "aepdf_rpt@example.com", "AEMP_EPDF", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)

        res = client.get("/reports/approvals/export/pdf", headers=headers)
        assert res.status_code == 200

    def test_filter_by_status(self, client, db_session, make_token):
        user = _create_user(db_session, "apdf2_rpt@example.com", "AEMP_PDF2", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers).json()
        _create_audit_log(db_session, user.id, "approve", "decision", r["id"])
        _create_audit_log(db_session, user.id, "reject", "decision", r["id"])

        res = client.get("/reports/approvals/export/pdf?status=approved", headers=headers)
        assert res.status_code == 200

    def test_invalid_status_returns_422(self, client, db_session, make_token):
        user = _create_user(db_session, "apdf3_rpt@example.com", "AEMP_PDF3", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/approvals/export/pdf?status=invalid", headers=headers)
        assert res.status_code == 422

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/approvals/export/pdf")
        assert res.status_code == 401


class TestExportTeamsPDF:
    def test_returns_pdf(self, client, db_session, make_token):
        _create_user(db_session, "tpdf_rpt@example.com", "TEMP_PDF", UserRole.EMPLOYEE, department="Engineering")
        admin = _create_user(db_session, "tpdfadm_rpt@example.com", "TEMP_PDFADM", UserRole.ADMINISTRATOR)
        headers = _auth_headers(admin, make_token)

        res = client.get("/reports/teams/export/pdf", headers=headers)
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "tepdf_rpt@example.com", "TEMP_EPDF", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)

        res = client.get("/reports/teams/export/pdf", headers=headers)
        assert res.status_code == 200

    def test_filter_by_team(self, client, db_session, make_token):
        _create_user(db_session, "tpdf2_rpt@example.com", "TEMP_PDF2", UserRole.EMPLOYEE, department="Engineering")
        _create_user(db_session, "tpdf3_rpt@example.com", "TEMP_PDF3", UserRole.EMPLOYEE, department="Marketing")
        admin = _create_user(db_session, "tpdf4_rpt@example.com", "TEMP_PDF4", UserRole.ADMINISTRATOR)
        headers = _auth_headers(admin, make_token)

        res = client.get("/reports/teams/export/pdf?team=Engineering", headers=headers)
        assert res.status_code == 200

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/teams/export/pdf")
        assert res.status_code == 401


class TestExportAuditPDF:
    def test_returns_pdf(self, client, db_session, make_token):
        user = _create_user(db_session, "updf_rpt@example.com", "UEMP_PDF", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_audit_log(db_session, user.id, "create", "decision", 1, "Created", "10.0.0.1")

        res = client.get("/reports/audit/export/pdf", headers=headers)
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "uepdf_rpt@example.com", "UEMP_EPDF", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)

        res = client.get("/reports/audit/export/pdf", headers=headers)
        assert res.status_code == 200

    def test_filter_by_action(self, client, db_session, make_token):
        user = _create_user(db_session, "updf2_rpt@example.com", "UEMP_PDF2", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_audit_log(db_session, user.id, "create", "decision", 1)
        _create_audit_log(db_session, user.id, "update", "decision", 1)

        res = client.get("/reports/audit/export/pdf?action=create", headers=headers)
        assert res.status_code == 200

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/audit/export/pdf")
        assert res.status_code == 401


# ===========================================================================
# Excel Export Tests
# ===========================================================================

class TestExportDecisionsExcel:
    def test_returns_excel(self, client, db_session, make_token):
        user = _create_user(db_session, "dexcel_rpt@example.com", "DEMP_EXC", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_decision(client, headers, title="Decision A")

        res = client.get("/reports/decisions/export/excel", headers=headers)
        assert res.status_code == 200
        assert "spreadsheetml" in res.headers["content-type"]
        assert "decisions_report.xlsx" in res.headers.get("content-disposition", "")

    def test_valid_xlsx(self, client, db_session, make_token):
        user = _create_user(db_session, "dexc2_rpt@example.com", "DEMP_EXC2", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_decision(client, headers, title="Alpha")
        _create_decision(client, headers, title="Beta")

        res = client.get("/reports/decisions/export/excel", headers=headers)
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        # Title row + timestamp + summary rows + header + data
        assert ws.cell(row=1, column=1).value == "Decisions Report"
        # Find header row (row with "ID" in column 1)
        header_row = None
        for row in range(1, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == "ID":
                header_row = row
                break
        assert header_row is not None
        headers = [ws.cell(row=header_row, column=c).value for c in range(1, 10)]
        assert "Title" in headers
        assert "Status" in headers

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "dempexc_rpt@example.com", "DEMP_EEXC", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)

        res = client.get("/reports/decisions/export/excel", headers=headers)
        assert res.status_code == 200
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Decisions Report"

    def test_filter_by_status(self, client, db_session, make_token):
        user = _create_user(db_session, "dexc3_rpt@example.com", "DEMP_EXC3", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers, title="Approved Dec").json()
        client.patch(f"/decisions/{r['id']}/status", json={"status": "Approved"}, headers=headers)
        _create_decision(client, headers, title="Draft Dec")

        res = client.get("/reports/decisions/export/excel?status=Approved", headers=headers)
        assert res.status_code == 200
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        # Find data rows
        data_rows = []
        for row in range(1, ws.max_row + 1):
            if ws.cell(row=row, column=1).value and str(ws.cell(row=row, column=1).value).isdigit():
                data_rows.append(row)
        # Only approved decisions should appear
        for r in data_rows:
            status_col = headers.index("Status") + 1 if "Status" in headers else 4
            assert ws.cell(row=r, column=status_col).value == "Approved"

    def test_invalid_date_returns_422(self, client, db_session, make_token):
        user = _create_user(db_session, "dexc4_rpt@example.com", "DEMP_EXC4", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/decisions/export/excel?start_date=bad", headers=headers)
        assert res.status_code == 422

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/decisions/export/excel")
        assert res.status_code == 401


class TestExportApprovalsExcel:
    def test_returns_excel(self, client, db_session, make_token):
        user = _create_user(db_session, "aexcel_rpt@example.com", "AEMP_EXC", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers).json()
        _create_audit_log(db_session, user.id, "approve", "decision", r["id"], "Approved")

        res = client.get("/reports/approvals/export/excel", headers=headers)
        assert res.status_code == 200
        assert "spreadsheetml" in res.headers["content-type"]

    def test_valid_xlsx(self, client, db_session, make_token):
        user = _create_user(db_session, "aexc2_rpt@example.com", "AEMP_EXC2", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers).json()
        _create_audit_log(db_session, user.id, "approve", "decision", r["id"])

        res = client.get("/reports/approvals/export/excel", headers=headers)
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Approvals Report"

    def test_filter_by_status(self, client, db_session, make_token):
        user = _create_user(db_session, "aexc3_rpt@example.com", "AEMP_EXC3", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers).json()
        _create_audit_log(db_session, user.id, "approve", "decision", r["id"])
        _create_audit_log(db_session, user.id, "reject", "decision", r["id"])

        res = client.get("/reports/approvals/export/excel?status=approved", headers=headers)
        assert res.status_code == 200

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "aempexc_rpt@example.com", "AEMP_EEXC", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/approvals/export/excel", headers=headers)
        assert res.status_code == 200

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/approvals/export/excel")
        assert res.status_code == 401


class TestExportTeamsExcel:
    def test_returns_excel(self, client, db_session, make_token):
        _create_user(db_session, "texcel_rpt@example.com", "TEMP_EXC", UserRole.EMPLOYEE, department="Engineering")
        admin = _create_user(db_session, "texceladm_rpt@example.com", "TEMP_EXCADM", UserRole.ADMINISTRATOR)
        headers = _auth_headers(admin, make_token)

        res = client.get("/reports/teams/export/excel", headers=headers)
        assert res.status_code == 200
        assert "spreadsheetml" in res.headers["content-type"]

    def test_valid_xlsx(self, client, db_session, make_token):
        _create_user(db_session, "texc2_rpt@example.com", "TEMP_EXC2", UserRole.EMPLOYEE, department="Engineering")
        admin = _create_user(db_session, "texc3_rpt@example.com", "TEMP_EXC3", UserRole.ADMINISTRATOR)
        headers = _auth_headers(admin, make_token)

        res = client.get("/reports/teams/export/excel", headers=headers)
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Teams Report"

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "tempexc_rpt@example.com", "TEMP_EEXC", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/teams/export/excel", headers=headers)
        assert res.status_code == 200

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/teams/export/excel")
        assert res.status_code == 401


class TestExportAuditExcel:
    def test_returns_excel(self, client, db_session, make_token):
        user = _create_user(db_session, "uexcel_rpt@example.com", "UEMP_EXC", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_audit_log(db_session, user.id, "create", "decision", 1, "Created", "10.0.0.1")

        res = client.get("/reports/audit/export/excel", headers=headers)
        assert res.status_code == 200
        assert "spreadsheetml" in res.headers["content-type"]

    def test_valid_xlsx(self, client, db_session, make_token):
        user = _create_user(db_session, "uexc2_rpt@example.com", "UEMP_EXC2", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_audit_log(db_session, user.id, "create", "decision", 1)

        res = client.get("/reports/audit/export/excel", headers=headers)
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Audit Report"

    def test_filter_by_action(self, client, db_session, make_token):
        user = _create_user(db_session, "uexc3_rpt@example.com", "UEMP_EXC3", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_audit_log(db_session, user.id, "create", "decision", 1)
        _create_audit_log(db_session, user.id, "update", "decision", 1)

        res = client.get("/reports/audit/export/excel?action=create", headers=headers)
        assert res.status_code == 200

    def test_empty_result(self, client, db_session, make_token):
        user = _create_user(db_session, "uempexc_rpt@example.com", "UEMP_EEXC", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        res = client.get("/reports/audit/export/excel", headers=headers)
        assert res.status_code == 200

    def test_requires_auth(self, client, db_session):
        res = client.get("/reports/audit/export/excel")
        assert res.status_code == 401


# ===========================================================================
# Integration: verify filter actually filters data in exports
# ===========================================================================

class TestExportFilterIntegration:
    def test_decisions_pdf_filter_actual_data(self, client, db_session, make_token):
        user = _create_user(db_session, "ifilt_rpt@example.com", "IEMP_FILT", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r1 = _create_decision(client, headers, title="Draft One", category="Engineering").json()
        r2 = _create_decision(client, headers, title="Draft Two", category="Marketing").json()
        client.patch(f"/decisions/{r1['id']}/status", json={"status": "Approved"}, headers=headers)

        # PDF with status=Approved — verify valid PDF and non-empty
        res = client.get("/reports/decisions/export/pdf?status=Approved", headers=headers)
        assert res.status_code == 200
        assert res.content[:5] == b"%PDF-"
        assert len(res.content) > 200  # non-trivial size

        # PDF with status=Draft — different size/content (only Draft Two)
        res_draft = client.get("/reports/decisions/export/pdf?status=Draft", headers=headers)
        assert res_draft.status_code == 200
        # The two PDFs should differ in content (different data)
        assert res.content != res_draft.content

    def test_decisions_excel_filter_actual_data(self, client, db_session, make_token):
        user = _create_user(db_session, "ifilt2_rpt@example.com", "IEMP_FILT2", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_decision(client, headers, title="Eng Only", category="Engineering")
        _create_decision(client, headers, title="Mkt Only", category="Marketing")

        res = client.get("/reports/decisions/export/excel?category=Engineering", headers=headers)
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        found_titles = []
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=2).value
            if val and val not in ("Title", "Decisions Report", None):
                found_titles.append(val)
        assert "Eng Only" in found_titles
        assert "Mkt Only" not in found_titles

    def test_approvals_pdf_only_approved(self, client, db_session, make_token):
        user = _create_user(db_session, "aifilt_rpt@example.com", "AIEMP_FILT", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        r = _create_decision(client, headers).json()
        _create_audit_log(db_session, user.id, "approve", "decision", r["id"])
        _create_audit_log(db_session, user.id, "reject", "decision", r["id"])

        # approved PDF vs rejected PDF — should produce different content
        res_approved = client.get("/reports/approvals/export/pdf?status=approved", headers=headers)
        res_rejected = client.get("/reports/approvals/export/pdf?status=rejected", headers=headers)
        assert res_approved.status_code == 200
        assert res_rejected.status_code == 200
        assert res_approved.content[:5] == b"%PDF-"
        # Different data = different PDF content
        assert res_approved.content != res_rejected.content

    def test_audit_excel_filter_by_entity(self, client, db_session, make_token):
        user = _create_user(db_session, "uifilt_rpt@example.com", "UIEMP_FILT", UserRole.ADMINISTRATOR)
        headers = _auth_headers(user, make_token)
        _create_audit_log(db_session, user.id, "create", "decision", 1)
        _create_audit_log(db_session, user.id, "create", "comment", 2)

        res = client.get("/reports/audit/export/excel?entity_type=decision", headers=headers)
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        entity_types = []
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=4).value
            if val and val not in ("Entity Type", "Audit Report", None, ""):
                entity_types.append(val)
        assert all(et == "decision" for et in entity_types)
        assert len(entity_types) >= 1
