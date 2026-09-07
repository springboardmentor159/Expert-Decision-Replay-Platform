from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, Query, HTTPException, status as http_status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from app.db.database import get_db
from app.models.decision import Decision
from app.models.user import User
from app.models.tag import Tag
from app.models.approval import Approval
from app.models.audit_log import AuditLog
from app.core.dependencies import get_current_user


router = APIRouter(
    prefix="/reports",
    tags=["Reports"]
)


# =========================================================
# ROLE HELPER
# =========================================================

def get_role(user):
    return (
        user.role.value
        if hasattr(user.role, "value")
        else user.role
    )


# =========================================================
# CONSTANTS
# =========================================================

ALLOWED_DECISION_STATUSES = {
    "Draft",
    "Under Review",
    "Approved",
    "Rejected",
    "Archived"
}

ALLOWED_APPROVAL_STATUSES = {
    "Pending",
    "Approved",
    "Rejected"
}

ALLOWED_AUDIT_ACTIONS = {
    "CREATE",
    "UPDATE",
    "DELETE",
    "APPROVE",
    "REJECT",
    "SUBMIT",
    "LOGIN",
    "LOGOUT",
    "ACCESS",
    "ARCHIVE"
}

ALLOWED_AUDIT_ENTITY_TYPES = {
    "Decision",
    "Alternative",
    "Comment",
    "DiscussionThread",
    "MeetingNote",
    "Approval",
    "User",
    "AuditLog",
    "SecurityLog",
    "AccessLog"
}


# =========================================================
# COMMON VALIDATION
# =========================================================

def validate_pagination(page: int, page_size: int):
    if page < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="page must be greater than or equal to 1"
        )

    if page_size < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="page_size must be greater than or equal to 1"
        )

    if page_size > 100:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="page_size must be less than or equal to 100"
        )


def validate_date_range(
    start_date: datetime | None,
    end_date: datetime | None
):
    if start_date and end_date and start_date > end_date:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="start_date cannot be greater than end_date"
        )


def validate_sort(
    sort_by: str,
    sort_order: str,
    allowed_fields: set
):
    if sort_by not in allowed_fields:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Invalid sort_by. Allowed values are: "
                + ", ".join(sorted(allowed_fields))
            )
        )

    if sort_order not in {"asc", "desc"}:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="sort_order must be either asc or desc"
        )


def require_admin(current_user: User):
    if get_role(current_user) != "Administrator":
        raise HTTPException(
            status_code=http_status.HTTP_403_FORBIDDEN,
            detail="Only Administrator can access audit reports"
        )


# =========================================================
# DECISION REPORT DATA HELPER
# =========================================================

def get_decision_report_data(
    db: Session,
    category: str | None = None,
    decision_status: str | None = None,
    created_by: int | None = None,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
    tag: str | None = None,
    sort_by: str = "created_at",
    sort_order: str = "desc"
):

    query = (
        db.query(Decision)
        .join(User, Decision.created_by == User.id)
    )

    if category:
        query = query.filter(
            Decision.category == category
        )

    if decision_status:
        query = query.filter(
            Decision.status == decision_status
        )

    if created_by:
        query = query.filter(
            Decision.created_by == created_by
        )

    if start_date:
        query = query.filter(
            Decision.created_at >= start_date
        )

    if end_date:
        query = query.filter(
            Decision.created_at <= end_date
        )

    if tag:
        query = (
            query
            .join(Decision.tags)
            .filter(Tag.name == tag)
        )

    query = query.distinct()

    sort_columns = {
        "id": Decision.id,
        "title": Decision.title,
        "category": Decision.category,
        "status": Decision.status,
        "created_at": Decision.created_at,
        "updated_at": Decision.updated_at
    }

    if sort_order == "asc":
        query = query.order_by(
            sort_columns[sort_by].asc()
        )
    else:
        query = query.order_by(
            sort_columns[sort_by].desc()
        )

    decisions = query.all()

    return decisions


def build_decision_rows(decisions):

    rows = []

    for decision in decisions:

        rows.append({
            "decision_id": decision.id,
            "title": decision.title,
            "category": decision.category,
            "status": decision.status,
            "created_by_id": decision.user.id,
            "created_by": decision.user.full_name,
            "created_date": decision.created_at,
            "updated_date": decision.updated_at,
            "number_of_alternatives": len(
                decision.alternatives
            ),
            "number_of_approvals": len(
                decision.approvals
            ),
            "tags": ", ".join(
                tag.name
                for tag in decision.tags
            )
        })

    return rows


def build_decision_summary(decisions):

    return {
        "total": len(decisions),
        "draft": sum(
            1 for d in decisions
            if d.status == "Draft"
        ),
        "under_review": sum(
            1 for d in decisions
            if d.status == "Under Review"
        ),
        "approved": sum(
            1 for d in decisions
            if d.status == "Approved"
        ),
        "rejected": sum(
            1 for d in decisions
            if d.status == "Rejected"
        ),
        "archived": sum(
            1 for d in decisions
            if d.status == "Archived"
        )
    }


# =========================================================
# DECISION REPORT
# =========================================================

@router.get("/decisions")
def get_decision_report(
    category: str | None = Query(default=None),
    status: str | None = Query(default=None),
    created_by: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    tag: str | None = Query(default=None),
    page: int = Query(default=1),
    page_size: int = Query(default=10),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    validate_pagination(page, page_size)

    if created_by is not None and created_by < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="created_by must be greater than or equal to 1"
        )

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "title",
            "category",
            "status",
            "created_at",
            "updated_at"
        }
    )

    if status and status not in ALLOWED_DECISION_STATUSES:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Invalid status. Allowed values are: "
                "Draft, Under Review, Approved, Rejected, Archived"
            )
        )

    validate_date_range(
        start_date,
        end_date
    )

    decisions = get_decision_report_data(
        db=db,
        category=category,
        decision_status=status,
        created_by=created_by,
        start_date=start_date,
        end_date=end_date,
        tag=tag,
        sort_by=sort_by,
        sort_order=sort_order
    )

    total = len(decisions)

    summary = build_decision_summary(
        decisions
    )

    offset = (page - 1) * page_size

    paginated_decisions = decisions[
        offset:
        offset + page_size
    ]

    report_data = build_decision_rows(
        paginated_decisions
    )

    return {
        "report": "Decision Report",

        "filters": {
            "category": category,
            "status": status,
            "created_by": created_by,
            "start_date": start_date,
            "end_date": end_date,
            "tag": tag
        },

        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_records": total,
            "total_pages": (
                (total + page_size - 1) // page_size
                if total > 0
                else 0
            )
        },

        "summary": summary,

        "data": report_data
    }


# =========================================================
# APPROVAL REPORT DATA HELPER
# =========================================================

def get_approval_report_data(
    db: Session,
    approval_status: str | None = None,
    reviewer: int | None = None,
    decision: int | None = None,
    level: int | None = None,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
    sort_by: str = "created_at",
    sort_order: str = "desc"
):

    query = (
        db.query(Approval)
        .join(
            Decision,
            Approval.decision_id == Decision.id
        )
        .join(
            User,
            Approval.reviewer_id == User.id
        )
    )

    if approval_status:
        query = query.filter(
            Approval.status == approval_status
        )

    if reviewer:
        query = query.filter(
            Approval.reviewer_id == reviewer
        )

    if decision:
        query = query.filter(
            Approval.decision_id == decision
        )

    if level:
        query = query.filter(
            Approval.approval_level == level
        )

    if start_date:
        query = query.filter(
            Approval.created_at >= start_date
        )

    if end_date:
        query = query.filter(
            Approval.created_at <= end_date
        )

    sort_columns = {
        "id": Approval.id,
        "decision_id": Approval.decision_id,
        "reviewer_id": Approval.reviewer_id,
        "approval_level": Approval.approval_level,
        "status": Approval.status,
        "created_at": Approval.created_at,
        "completed_at": Approval.completed_at
    }

    if sort_order == "asc":
        query = query.order_by(
            sort_columns[sort_by].asc()
        )
    else:
        query = query.order_by(
            sort_columns[sort_by].desc()
        )

    return query.all()


def build_approval_rows(approvals):

    rows = []

    for approval in approvals:

        turnaround_seconds = None
        turnaround_time = None

        if approval.completed_at:

            difference = (
                approval.completed_at -
                approval.created_at
            )

            turnaround_seconds = (
                difference.total_seconds()
            )

            days = difference.days
            hours = difference.seconds // 3600
            minutes = (
                difference.seconds % 3600
            ) // 60
            seconds = difference.seconds % 60

            turnaround_time = (
                f"{days} days, "
                f"{hours:02d}:"
                f"{minutes:02d}:"
                f"{seconds:02d}"
            )

        rows.append({
            "approval_id": approval.id,
            "decision_id": approval.decision.id,
            "decision_title": approval.decision.title,
            "reviewer_id": approval.reviewer.id,
            "reviewer": approval.reviewer.full_name,
            "approval_level": approval.approval_level,
            "status": approval.status,
            "assigned_date": approval.created_at,
            "completed_date": approval.completed_at,
            "turnaround_time": turnaround_time,
            "turnaround_seconds": turnaround_seconds
        })

    return rows


def build_approval_summary(approvals):

    total = len(approvals)

    pending = sum(
        1 for a in approvals
        if a.status == "Pending"
    )

    approved = sum(
        1 for a in approvals
        if a.status == "Approved"
    )

    rejected = sum(
        1 for a in approvals
        if a.status == "Rejected"
    )

    turnaround_values = []

    for approval in approvals:

        if approval.completed_at:

            difference = (
                approval.completed_at -
                approval.created_at
            )

            turnaround_values.append(
                difference.total_seconds()
            )

    average_turnaround = (
        sum(turnaround_values)
        / len(turnaround_values)
        if turnaround_values
        else 0
    )

    completed = approved + rejected

    completion_rate = (
        (completed / total) * 100
        if total > 0
        else 0
    )

    return {
        "total": total,
        "pending": pending,
        "approved": approved,
        "rejected": rejected,
        "avg_turnaround_seconds": round(
            average_turnaround,
            2
        ),
        "completion_rate": round(
            completion_rate,
            2
        )
    }


# =========================================================
# APPROVAL REPORT
# =========================================================

@router.get("/approvals")
def get_approval_report(
    status: str | None = Query(default=None),
    reviewer: int | None = Query(default=None),
    decision: int | None = Query(default=None),
    level: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    page: int = Query(default=1),
    page_size: int = Query(default=10),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    validate_pagination(
        page,
        page_size
    )

    if reviewer is not None and reviewer < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="reviewer must be greater than or equal to 1"
        )

    if decision is not None and decision < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="decision must be greater than or equal to 1"
        )

    if level is not None and level < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="level must be greater than or equal to 1"
        )

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "decision_id",
            "reviewer_id",
            "approval_level",
            "status",
            "created_at",
            "completed_at"
        }
    )

    if status and status not in ALLOWED_APPROVAL_STATUSES:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Invalid status. Allowed values are: "
                "Pending, Approved, Rejected"
            )
        )

    validate_date_range(
        start_date,
        end_date
    )

    approvals = get_approval_report_data(
        db=db,
        approval_status=status,
        reviewer=reviewer,
        decision=decision,
        level=level,
        start_date=start_date,
        end_date=end_date,
        sort_by=sort_by,
        sort_order=sort_order
    )

    total = len(approvals)

    offset = (page - 1) * page_size

    paginated_approvals = approvals[
        offset:
        offset + page_size
    ]

    return {
        "report": "Approval Report",

        "filters": {
            "status": status,
            "reviewer": reviewer,
            "decision": decision,
            "level": level,
            "start_date": start_date,
            "end_date": end_date
        },

        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_records": total,
            "total_pages": (
                (total + page_size - 1) // page_size
                if total > 0
                else 0
            )
        },

        "summary": build_approval_summary(
            approvals
        ),

        "data": build_approval_rows(
            paginated_approvals
        )
    }


# =========================================================
# TEAM REPORT
# =========================================================

@router.get("/teams")
def get_team_report(
    team: str | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    status: str | None = Query(default=None),
    category: str | None = Query(default=None),
    page: int = Query(default=1),
    page_size: int = Query(default=10),
    sort_by: str = Query(default="team_name"),
    sort_order: str = Query(default="asc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    validate_pagination(
        page,
        page_size
    )

    validate_sort(
        sort_by,
        sort_order,
        {
            "team_name",
            "member_count",
            "total_decisions",
            "approved_decisions",
            "rejected_decisions",
            "pending_decisions",
            "total_approvals",
            "approved_approvals",
            "rejected_approvals",
            "pending_approvals",
            "avg_turnaround_seconds"
        }
    )

    role = get_role(current_user)

    if role not in {
        "Manager",
        "Administrator"
    }:
        raise HTTPException(
            status_code=http_status.HTTP_403_FORBIDDEN,
            detail=(
                "Only Manager and Administrator "
                "can access team reports"
            )
        )

    if status and status not in ALLOWED_DECISION_STATUSES:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Invalid status. Allowed values are: "
                "Draft, Under Review, Approved, "
                "Rejected, Archived"
            )
        )

    validate_date_range(
        start_date,
        end_date
    )

    if role == "Manager":

        permitted_team = current_user.department

        if team and team != permitted_team:
            raise HTTPException(
                status_code=http_status.HTTP_403_FORBIDDEN,
                detail=(
                    "Manager can only access "
                    "their own team report"
                )
            )

        teams_query = (
            db.query(User.department)
            .filter(
                User.department == permitted_team
            )
            .distinct()
        )

    else:

        teams_query = (
            db.query(User.department)
            .filter(
                User.department.isnot(None)
            )
            .distinct()
        )

        if team:
            teams_query = teams_query.filter(
                User.department == team
            )

    team_rows = teams_query.all()

    team_names = [
        row[0]
        for row in team_rows
        if row[0]
    ]

    report_data = []

    for team_name in team_names:

        member_count = (
            db.query(User)
            .filter(
                User.department == team_name
            )
            .count()
        )

        decision_query = (
            db.query(Decision)
            .join(
                User,
                Decision.created_by == User.id
            )
            .filter(
                User.department == team_name
            )
        )

        if start_date:
            decision_query = decision_query.filter(
                Decision.created_at >= start_date
            )

        if end_date:
            decision_query = decision_query.filter(
                Decision.created_at <= end_date
            )

        if status:
            decision_query = decision_query.filter(
                Decision.status == status
            )

        if category:
            decision_query = decision_query.filter(
                Decision.category == category
            )

        filtered_decisions = (
            decision_query.all()
        )

        decision_ids = [
            decision.id
            for decision in filtered_decisions
        ]

        total_decisions = len(
            filtered_decisions
        )

        approved_decisions = sum(
            1
            for decision in filtered_decisions
            if decision.status == "Approved"
        )

        rejected_decisions = sum(
            1
            for decision in filtered_decisions
            if decision.status == "Rejected"
        )

        pending_decisions = sum(
            1
            for decision in filtered_decisions
            if decision.status in {
                "Draft",
                "Under Review"
            }
        )

        if decision_ids:

            approvals = (
                db.query(Approval)
                .filter(
                    Approval.decision_id.in_(
                        decision_ids
                    )
                )
                .all()
            )

        else:

            approvals = []

        total_approvals = len(
            approvals
        )

        approved_approvals = sum(
            1
            for approval in approvals
            if approval.status == "Approved"
        )

        rejected_approvals = sum(
            1
            for approval in approvals
            if approval.status == "Rejected"
        )

        pending_approvals = sum(
            1
            for approval in approvals
            if approval.status == "Pending"
        )

        turnaround_values = []

        for approval in approvals:

            if approval.completed_at:

                difference = (
                    approval.completed_at -
                    approval.created_at
                )

                turnaround_values.append(
                    difference.total_seconds()
                )

        avg_turnaround_seconds = (
            sum(turnaround_values)
            / len(turnaround_values)
            if turnaround_values
            else 0
        )

        completed_approvals = (
            approved_approvals +
            rejected_approvals
        )

        approval_completion_rate = (
            (
                completed_approvals /
                total_approvals
            ) * 100
            if total_approvals > 0
            else 0
        )

        report_data.append({

            "team_name": team_name,

            "member_count": member_count,

            "total_decisions": total_decisions,

            "approved_decisions": approved_decisions,

            "rejected_decisions": rejected_decisions,

            "pending_decisions": pending_decisions,

            "approval_statistics": {

                "total_approvals": total_approvals,

                "approved_approvals": approved_approvals,

                "rejected_approvals": rejected_approvals,

                "pending_approvals": pending_approvals,

                "avg_turnaround_seconds": round(
                    avg_turnaround_seconds,
                    2
                ),

                "completion_rate": round(
                    approval_completion_rate,
                    2
                )
            },

            "total_approvals": total_approvals,

            "approved_approvals": approved_approvals,

            "rejected_approvals": rejected_approvals,

            "pending_approvals": pending_approvals,

            "avg_turnaround_seconds": round(
                avg_turnaround_seconds,
                2
            )
        })

    sort_key_map = {

        "team_name": "team_name",

        "member_count": "member_count",

        "total_decisions": "total_decisions",

        "approved_decisions": "approved_decisions",

        "rejected_decisions": "rejected_decisions",

        "pending_decisions": "pending_decisions",

        "total_approvals": "total_approvals",

        "approved_approvals": "approved_approvals",

        "rejected_approvals": "rejected_approvals",

        "pending_approvals": "pending_approvals",

        "avg_turnaround_seconds":
            "avg_turnaround_seconds"
    }

    sort_key = sort_key_map[sort_by]

    report_data.sort(
        key=lambda item: (
            item[sort_key]
            if item[sort_key] is not None
            else 0
        ),
        reverse=(
            sort_order == "desc"
        )
    )

    total_teams = len(
        report_data
    )

    total_members = sum(
        item["member_count"]
        for item in report_data
    )

    total_decisions = sum(
        item["total_decisions"]
        for item in report_data
    )

    total_approved_decisions = sum(
        item["approved_decisions"]
        for item in report_data
    )

    total_rejected_decisions = sum(
        item["rejected_decisions"]
        for item in report_data
    )

    total_pending_decisions = sum(
        item["pending_decisions"]
        for item in report_data
    )

    total_approvals = sum(
        item["total_approvals"]
        for item in report_data
    )

    total_approved_approvals = sum(
        item["approved_approvals"]
        for item in report_data
    )

    total_rejected_approvals = sum(
        item["rejected_approvals"]
        for item in report_data
    )

    total_pending_approvals = sum(
        item["pending_approvals"]
        for item in report_data
    )

    turnaround_values = [
        item["avg_turnaround_seconds"]
        for item in report_data
        if item["avg_turnaround_seconds"] > 0
    ]

    overall_avg_turnaround = (
        sum(turnaround_values)
        / len(turnaround_values)
        if turnaround_values
        else 0
    )

    completed_total = (
        total_approved_approvals +
        total_rejected_approvals
    )

    overall_completion_rate = (
        (
            completed_total /
            total_approvals
        ) * 100
        if total_approvals > 0
        else 0
    )

    total_records = len(
        report_data
    )

    offset = (
        (page - 1) *
        page_size
    )

    paginated_data = report_data[
        offset:
        offset + page_size
    ]

    return {

        "report": "Team Report",

        "filters": {
            "team": team,
            "start_date": start_date,
            "end_date": end_date,
            "status": status,
            "category": category
        },

        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_records": total_records,
            "total_pages": (
                (
                    total_records +
                    page_size -
                    1
                ) // page_size
                if total_records > 0
                else 0
            )
        },

        "summary": {

            "total_teams": total_teams,

            "total_members": total_members,

            "total_decisions": total_decisions,

            "approved_decisions":
                total_approved_decisions,

            "rejected_decisions":
                total_rejected_decisions,

            "pending_decisions":
                total_pending_decisions,

            "approval_statistics": {

                "total_approvals":
                    total_approvals,

                "approved_approvals":
                    total_approved_approvals,

                "rejected_approvals":
                    total_rejected_approvals,

                "pending_approvals":
                    total_pending_approvals,

                "avg_turnaround_seconds": round(
                    overall_avg_turnaround,
                    2
                ),

                "completion_rate": round(
                    overall_completion_rate,
                    2
                )
            }
        },

        "data": paginated_data
    }


# =========================================================
# AUDIT REPORT
# =========================================================

@router.get("/audit")
def get_audit_report(
    user: int | None = Query(default=None),
    action: str | None = Query(default=None),
    entity_type: str | None = Query(default=None),
    entity_id: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    page: int = Query(default=1),
    page_size: int = Query(default=10),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    require_admin(current_user)

    validate_pagination(
        page,
        page_size
    )

    if user is not None and user < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="user must be greater than or equal to 1"
        )

    if entity_id is not None and entity_id < 1:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="entity_id must be greater than or equal to 1"
        )

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "user_id",
            "action",
            "entity_type",
            "entity_id",
            "created_at"
        }
    )

    if action and action not in ALLOWED_AUDIT_ACTIONS:
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Invalid action. Allowed values are: "
                + ", ".join(sorted(ALLOWED_AUDIT_ACTIONS))
            )
        )

    if (
        entity_type
        and entity_type not in ALLOWED_AUDIT_ENTITY_TYPES
    ):
        raise HTTPException(
            status_code=http_status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Invalid entity_type. Allowed values are: "
                + ", ".join(sorted(ALLOWED_AUDIT_ENTITY_TYPES))
            )
        )

    validate_date_range(
        start_date,
        end_date
    )

    query = (
        db.query(AuditLog)
        .join(
            User,
            AuditLog.user_id == User.id
        )
    )

    if user:
        query = query.filter(
            AuditLog.user_id == user
        )

    if action:
        query = query.filter(
            AuditLog.action == action
        )

    if entity_type:
        query = query.filter(
            AuditLog.entity_type == entity_type
        )

    if entity_id:
        query = query.filter(
            AuditLog.entity_id == entity_id
        )

    if start_date:
        query = query.filter(
            AuditLog.created_at >= start_date
        )

    if end_date:
        query = query.filter(
            AuditLog.created_at <= end_date
        )

    sort_columns = {

        "id": AuditLog.id,

        "user_id": AuditLog.user_id,

        "action": AuditLog.action,

        "entity_type": AuditLog.entity_type,

        "entity_id": AuditLog.entity_id,

        "created_at": AuditLog.created_at
    }

    if sort_order == "asc":
        query = query.order_by(
            sort_columns[sort_by].asc()
        )
    else:
        query = query.order_by(
            sort_columns[sort_by].desc()
        )

    total = query.count()

    offset = (
        (page - 1) *
        page_size
    )

    logs = (
        query
        .offset(offset)
        .limit(page_size)
        .all()
    )

    report_data = []

    for log in logs:

        report_data.append({

            "user": {
                "id": log.user.id,
                "name": log.user.full_name
            },

            "action": log.action,

            "entity_type": log.entity_type,

            "entity_id": log.entity_id,

            "description": log.description,

            "timestamp": log.created_at,

            "ip_address": log.ip_address
        })

    return {

        "report": "Audit Report",

        "filters": {

            "user": user,

            "action": action,

            "entity_type": entity_type,

            "entity_id": entity_id,

            "start_date": start_date,

            "end_date": end_date
        },

        "pagination": {

            "page": page,

            "page_size": page_size,

            "total_records": total,

            "total_pages": (
                (
                    total +
                    page_size -
                    1
                ) // page_size
                if total > 0
                else 0
            )
        },

        "summary": {
            "total": total
        },

        "data": report_data
    }


# =========================================================
# PDF HELPER
# =========================================================

def create_pdf(
    title,
    filters,
    summary,
    headers,
    rows
):

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            title,
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            "Generated Date: "
            + datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            8
        )
    )

    elements.append(
        Paragraph(
            "<b>Applied Filters</b>",
            styles["Heading3"]
        )
    )

    filter_text = []

    for key, value in filters.items():

        if value is not None:

            filter_text.append(
                f"{key}: {value}"
            )

    if not filter_text:
        filter_text.append(
            "No filters applied"
        )

    elements.append(
        Paragraph(
            "<br/>".join(filter_text),
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            8
        )
    )

    elements.append(
        Paragraph(
            "<b>Summary</b>",
            styles["Heading3"]
        )
    )

    summary_rows = [
        ["Metric", "Value"]
    ]

    for key, value in summary.items():

        summary_rows.append([
            str(key),
            str(value)
        ])

    summary_table = Table(
        summary_rows,
        repeatRows=1
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.grey
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.black
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "LEFT"
            )
        ])
    )

    elements.append(
        summary_table
    )

    elements.append(
        Spacer(
            1,
            10
        )
    )

    elements.append(
        Paragraph(
            "<b>Report Data</b>",
            styles["Heading3"]
        )
    )

    table_data = [
        headers
    ]

    for row in rows:

        table_data.append([
            str(
                value
                if value is not None
                else ""
            )
            for value in row
        ])

    if len(table_data) == 1:

        table_data.append([
            "No matching records"
        ] + [""] * (
            len(headers) - 1
        ))

    data_table = Table(
        table_data,
        repeatRows=1
    )

    data_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.grey
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.black
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP"
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7
            )
        ])
    )

    elements.append(
        data_table
    )

    document.build(
        elements
    )

    buffer.seek(0)

    return buffer


# =========================================================
# EXCEL HELPER
# =========================================================

def create_excel(
    title,
    filters,
    summary,
    headers,
    rows
):

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Report"

    worksheet["A1"] = title
    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = "Generated Date"
    worksheet["B2"] = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    worksheet["A4"] = "Applied Filters"
    worksheet["A4"].font = Font(
        bold=True
    )

    filter_row = 5

    for key, value in filters.items():

        worksheet.cell(
            row=filter_row,
            column=1,
            value=key
        )

        worksheet.cell(
            row=filter_row,
            column=2,
            value=(
                str(value)
                if value is not None
                else ""
            )
        )

        filter_row += 1

    summary_start = filter_row + 1

    worksheet.cell(
        row=summary_start,
        column=1,
        value="Summary"
    )

    worksheet.cell(
        row=summary_start,
        column=1
    ).font = Font(
        bold=True
    )

    summary_row = summary_start + 1

    for key, value in summary.items():

        worksheet.cell(
            row=summary_row,
            column=1,
            value=key
        )

        worksheet.cell(
            row=summary_row,
            column=2,
            value=value
        )

        summary_row += 1

    data_start = summary_row + 2

    for column_index, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=data_start,
            column=column_index,
            value=header
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for row_index, row in enumerate(
        rows,
        start=data_start + 1
    ):

        for column_index, value in enumerate(
            row,
            start=1
        ):

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=(
                    str(value)
                    if value is not None
                    else ""
                )
            )

    for column in worksheet.columns:

        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40
        )

    buffer = BytesIO()

    workbook.save(
        buffer
    )

    buffer.seek(0)

    return buffer


# =========================================================
# DECISION PDF EXPORT
# =========================================================

@router.get("/decisions/export/pdf")
def export_decision_pdf(
    category: str | None = Query(default=None),
    status: str | None = Query(default=None),
    created_by: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    tag: str | None = Query(default=None),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    if created_by is not None and created_by < 1:
        raise HTTPException(
            status_code=422,
            detail="created_by must be greater than or equal to 1"
        )

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "title",
            "category",
            "status",
            "created_at",
            "updated_at"
        }
    )

    if status and status not in ALLOWED_DECISION_STATUSES:
        raise HTTPException(
            status_code=422,
            detail="Invalid status"
        )

    validate_date_range(
        start_date,
        end_date
    )

    decisions = get_decision_report_data(
        db,
        category,
        status,
        created_by,
        start_date,
        end_date,
        tag,
        sort_by,
        sort_order
    )

    rows = build_decision_rows(
        decisions
    )

    pdf_rows = []

    for row in rows:

        pdf_rows.append([
            row["decision_id"],
            row["title"],
            row["category"],
            row["status"],
            row["created_by"],
            row["created_date"],
            row["updated_date"],
            row["number_of_alternatives"],
            row["number_of_approvals"],
            row["tags"]
        ])

    filters = {
        "category": category,
        "status": status,
        "created_by": created_by,
        "start_date": start_date,
        "end_date": end_date,
        "tag": tag
    }

    pdf = create_pdf(
        "Decision Report",
        filters,
        build_decision_summary(decisions),
        [
            "Decision ID",
            "Title",
            "Category",
            "Status",
            "Created By",
            "Created Date",
            "Updated Date",
            "Alternatives",
            "Approvals",
            "Tags"
        ],
        pdf_rows
    )

    return StreamingResponse(
        pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                "attachment; filename=decision_report.pdf"
        }
    )


# =========================================================
# DECISION EXCEL EXPORT
# =========================================================

@router.get("/decisions/export/excel")
def export_decision_excel(
    category: str | None = Query(default=None),
    status: str | None = Query(default=None),
    created_by: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    tag: str | None = Query(default=None),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    if created_by is not None and created_by < 1:
        raise HTTPException(
            status_code=422,
            detail="created_by must be greater than or equal to 1"
        )

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "title",
            "category",
            "status",
            "created_at",
            "updated_at"
        }
    )

    if status and status not in ALLOWED_DECISION_STATUSES:
        raise HTTPException(
            status_code=422,
            detail="Invalid status"
        )

    validate_date_range(
        start_date,
        end_date
    )

    decisions = get_decision_report_data(
        db,
        category,
        status,
        created_by,
        start_date,
        end_date,
        tag,
        sort_by,
        sort_order
    )

    rows = build_decision_rows(
        decisions
    )

    excel_rows = []

    for row in rows:

        excel_rows.append([
            row["decision_id"],
            row["title"],
            row["category"],
            row["status"],
            row["created_by"],
            row["created_date"],
            row["updated_date"],
            row["number_of_alternatives"],
            row["number_of_approvals"],
            row["tags"]
        ])

    excel = create_excel(
        "Decision Report",
        {
            "category": category,
            "status": status,
            "created_by": created_by,
            "start_date": start_date,
            "end_date": end_date,
            "tag": tag
        },
        build_decision_summary(decisions),
        [
            "Decision ID",
            "Title",
            "Category",
            "Status",
            "Created By",
            "Created Date",
            "Updated Date",
            "Number of Alternatives",
            "Number of Approvals",
            "Tags"
        ],
        excel_rows
    )

    return StreamingResponse(
        excel,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                "attachment; filename=decision_report.xlsx"
        }
    )


# =========================================================
# APPROVAL PDF EXPORT
# =========================================================

@router.get("/approvals/export/pdf")
def export_approval_pdf(
    status: str | None = Query(default=None),
    reviewer: int | None = Query(default=None),
    decision: int | None = Query(default=None),
    level: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "decision_id",
            "reviewer_id",
            "approval_level",
            "status",
            "created_at",
            "completed_at"
        }
    )

    if status and status not in ALLOWED_APPROVAL_STATUSES:
        raise HTTPException(
            status_code=422,
            detail="Invalid approval status"
        )

    validate_date_range(
        start_date,
        end_date
    )

    approvals = get_approval_report_data(
        db,
        status,
        reviewer,
        decision,
        level,
        start_date,
        end_date,
        sort_by,
        sort_order
    )

    rows = build_approval_rows(
        approvals
    )

    pdf_rows = []

    for row in rows:

        pdf_rows.append([
            row["approval_id"],
            row["decision_id"],
            row["decision_title"],
            row["reviewer"],
            row["approval_level"],
            row["status"],
            row["assigned_date"],
            row["completed_date"],
            row["turnaround_time"]
        ])

    pdf = create_pdf(
        "Approval Report",
        {
            "status": status,
            "reviewer": reviewer,
            "decision": decision,
            "level": level,
            "start_date": start_date,
            "end_date": end_date
        },
        build_approval_summary(
            approvals
        ),
        [
            "Approval ID",
            "Decision ID",
            "Decision Title",
            "Reviewer",
            "Level",
            "Status",
            "Assigned Date",
            "Completed Date",
            "Turnaround Time"
        ],
        pdf_rows
    )

    return StreamingResponse(
        pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                "attachment; filename=approval_report.pdf"
        }
    )


# =========================================================
# APPROVAL EXCEL EXPORT
# =========================================================

@router.get("/approvals/export/excel")
def export_approval_excel(
    status: str | None = Query(default=None),
    reviewer: int | None = Query(default=None),
    decision: int | None = Query(default=None),
    level: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "decision_id",
            "reviewer_id",
            "approval_level",
            "status",
            "created_at",
            "completed_at"
        }
    )

    if status and status not in ALLOWED_APPROVAL_STATUSES:
        raise HTTPException(
            status_code=422,
            detail="Invalid approval status"
        )

    validate_date_range(
        start_date,
        end_date
    )

    approvals = get_approval_report_data(
        db,
        status,
        reviewer,
        decision,
        level,
        start_date,
        end_date,
        sort_by,
        sort_order
    )

    rows = build_approval_rows(
        approvals
    )

    excel_rows = []

    for row in rows:

        excel_rows.append([
            row["approval_id"],
            row["decision_id"],
            row["decision_title"],
            row["reviewer"],
            row["approval_level"],
            row["status"],
            row["assigned_date"],
            row["completed_date"],
            row["turnaround_time"]
        ])

    excel = create_excel(
        "Approval Report",
        {
            "status": status,
            "reviewer": reviewer,
            "decision": decision,
            "level": level,
            "start_date": start_date,
            "end_date": end_date
        },
        build_approval_summary(
            approvals
        ),
        [
            "Approval ID",
            "Decision ID",
            "Decision Title",
            "Reviewer",
            "Approval Level",
            "Status",
            "Assigned Date",
            "Completed Date",
            "Turnaround Time"
        ],
        excel_rows
    )

    return StreamingResponse(
        excel,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                "attachment; filename=approval_report.xlsx"
        }
    )


# =========================================================
# TEAM PDF EXPORT
# =========================================================

@router.get("/teams/export/pdf")
def export_team_pdf(
    team: str | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    status: str | None = Query(default=None),
    category: str | None = Query(default=None),
    sort_by: str = Query(default="team_name"),
    sort_order: str = Query(default="asc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    role = get_role(current_user)

    if role not in {
        "Manager",
        "Administrator"
    }:
        raise HTTPException(
            status_code=403,
            detail=(
                "Only Manager and Administrator "
                "can access team reports"
            )
        )

    validate_sort(
        sort_by,
        sort_order,
        {
            "team_name",
            "member_count",
            "total_decisions",
            "approved_decisions",
            "rejected_decisions",
            "pending_decisions",
            "total_approvals",
            "approved_approvals",
            "rejected_approvals",
            "pending_approvals",
            "avg_turnaround_seconds"
        }
    )

    if role == "Manager":

        if team and team != current_user.department:
            raise HTTPException(
                status_code=403,
                detail=(
                    "Manager can only access "
                    "their own team report"
                )
            )

        selected_team = current_user.department

    else:

        selected_team = team

    if status and status not in ALLOWED_DECISION_STATUSES:
        raise HTTPException(
            status_code=422,
            detail="Invalid status"
        )

    validate_date_range(
        start_date,
        end_date
    )

    team_names = (
        [selected_team]
        if selected_team
        else [
            row[0]
            for row in (
                db.query(User.department)
                .filter(
                    User.department.isnot(None)
                )
                .distinct()
                .all()
            )
            if row[0]
        ]
    )

    team_data = []

    for team_name in team_names:

        member_count = (
            db.query(User)
            .filter(
                User.department == team_name
            )
            .count()
        )

        query = (
            db.query(Decision)
            .join(
                User,
                Decision.created_by == User.id
            )
            .filter(
                User.department == team_name
            )
        )

        if start_date:
            query = query.filter(
                Decision.created_at >= start_date
            )

        if end_date:
            query = query.filter(
                Decision.created_at <= end_date
            )

        if status:
            query = query.filter(
                Decision.status == status
            )

        if category:
            query = query.filter(
                Decision.category == category
            )

        decisions = query.all()

        decision_ids = [
            d.id
            for d in decisions
        ]

        approvals = (
            db.query(Approval)
            .filter(
                Approval.decision_id.in_(
                    decision_ids
                )
            )
            .all()
            if decision_ids
            else []
        )

        turnaround_values = []

        for approval in approvals:

            if approval.completed_at:

                turnaround_values.append(
                    (
                        approval.completed_at -
                        approval.created_at
                    ).total_seconds()
                )

        avg_turnaround = (
            sum(turnaround_values)
            / len(turnaround_values)
            if turnaround_values
            else 0
        )

        team_data.append({

            "team_name": team_name,

            "member_count": member_count,

            "total_decisions": len(decisions),

            "approved_decisions": sum(
                1 for d in decisions
                if d.status == "Approved"
            ),

            "rejected_decisions": sum(
                1 for d in decisions
                if d.status == "Rejected"
            ),

            "pending_decisions": sum(
                1 for d in decisions
                if d.status in {
                    "Draft",
                    "Under Review"
                }
            ),

            "total_approvals": len(
                approvals
            ),

            "approved_approvals": sum(
                1 for a in approvals
                if a.status == "Approved"
            ),

            "rejected_approvals": sum(
                1 for a in approvals
                if a.status == "Rejected"
            ),

            "pending_approvals": sum(
                1 for a in approvals
                if a.status == "Pending"
            ),

            "avg_turnaround_seconds": round(
                avg_turnaround,
                2
            )
        })

    team_data.sort(
        key=lambda x: (
            x[sort_by]
            if x[sort_by] is not None
            else 0
        ),
        reverse=(
            sort_order == "desc"
        )
    )

    pdf_rows = []

    for row in team_data:

        pdf_rows.append([
            row["team_name"],
            row["member_count"],
            row["total_decisions"],
            row["approved_decisions"],
            row["rejected_decisions"],
            row["pending_decisions"],
            row["total_approvals"],
            row["approved_approvals"],
            row["rejected_approvals"],
            row["pending_approvals"],
            row["avg_turnaround_seconds"]
        ])

    summary = {

        "total_teams": len(team_data),

        "total_members": sum(
            r["member_count"]
            for r in team_data
        ),

        "total_decisions": sum(
            r["total_decisions"]
            for r in team_data
        ),

        "approved_decisions": sum(
            r["approved_decisions"]
            for r in team_data
        ),

        "rejected_decisions": sum(
            r["rejected_decisions"]
            for r in team_data
        ),

        "pending_decisions": sum(
            r["pending_decisions"]
            for r in team_data
        ),

        "total_approvals": sum(
            r["total_approvals"]
            for r in team_data
        )
    }

    pdf = create_pdf(
        "Team Report",
        {
            "team": team,
            "start_date": start_date,
            "end_date": end_date,
            "status": status,
            "category": category
        },
        summary,
        [
            "Team",
            "Members",
            "Decisions",
            "Approved",
            "Rejected",
            "Pending",
            "Approvals",
            "Approved App.",
            "Rejected App.",
            "Pending App.",
            "Avg Turnaround"
        ],
        pdf_rows
    )

    return StreamingResponse(
        pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                "attachment; filename=team_report.pdf"
        }
    )


# =========================================================
# TEAM EXCEL EXPORT
# =========================================================

@router.get("/teams/export/excel")
def export_team_excel(
    team: str | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    status: str | None = Query(default=None),
    category: str | None = Query(default=None),
    sort_by: str = Query(default="team_name"),
    sort_order: str = Query(default="asc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    role = get_role(current_user)

    if role not in {
        "Manager",
        "Administrator"
    }:
        raise HTTPException(
            status_code=403,
            detail=(
                "Only Manager and Administrator "
                "can access team reports"
            )
        )

    if role == "Manager":

        if team and team != current_user.department:
            raise HTTPException(
                status_code=403,
                detail=(
                    "Manager can only access "
                    "their own team report"
                )
            )

        selected_team = current_user.department

    else:

        selected_team = team

    validate_sort(
        sort_by,
        sort_order,
        {
            "team_name",
            "member_count",
            "total_decisions",
            "approved_decisions",
            "rejected_decisions",
            "pending_decisions",
            "total_approvals",
            "approved_approvals",
            "rejected_approvals",
            "pending_approvals",
            "avg_turnaround_seconds"
        }
    )

    if status and status not in ALLOWED_DECISION_STATUSES:
        raise HTTPException(
            status_code=422,
            detail="Invalid status"
        )

    validate_date_range(
        start_date,
        end_date
    )

    team_names = (
        [selected_team]
        if selected_team
        else [
            row[0]
            for row in (
                db.query(User.department)
                .filter(
                    User.department.isnot(None)
                )
                .distinct()
                .all()
            )
            if row[0]
        ]
    )

    team_data = []

    for team_name in team_names:

        member_count = (
            db.query(User)
            .filter(
                User.department == team_name
            )
            .count()
        )

        query = (
            db.query(Decision)
            .join(
                User,
                Decision.created_by == User.id
            )
            .filter(
                User.department == team_name
            )
        )

        if start_date:
            query = query.filter(
                Decision.created_at >= start_date
            )

        if end_date:
            query = query.filter(
                Decision.created_at <= end_date
            )

        if status:
            query = query.filter(
                Decision.status == status
            )

        if category:
            query = query.filter(
                Decision.category == category
            )

        decisions = query.all()

        decision_ids = [
            d.id
            for d in decisions
        ]

        approvals = (
            db.query(Approval)
            .filter(
                Approval.decision_id.in_(
                    decision_ids
                )
            )
            .all()
            if decision_ids
            else []
        )

        turnaround_values = []

        for approval in approvals:

            if approval.completed_at:

                turnaround_values.append(
                    (
                        approval.completed_at -
                        approval.created_at
                    ).total_seconds()
                )

        avg_turnaround = (
            sum(turnaround_values)
            / len(turnaround_values)
            if turnaround_values
            else 0
        )

        team_data.append({

            "team_name": team_name,

            "member_count": member_count,

            "total_decisions": len(decisions),

            "approved_decisions": sum(
                1 for d in decisions
                if d.status == "Approved"
            ),

            "rejected_decisions": sum(
                1 for d in decisions
                if d.status == "Rejected"
            ),

            "pending_decisions": sum(
                1 for d in decisions
                if d.status in {
                    "Draft",
                    "Under Review"
                }
            ),

            "total_approvals": len(
                approvals
            ),

            "approved_approvals": sum(
                1 for a in approvals
                if a.status == "Approved"
            ),

            "rejected_approvals": sum(
                1 for a in approvals
                if a.status == "Rejected"
            ),

            "pending_approvals": sum(
                1 for a in approvals
                if a.status == "Pending"
            ),

            "avg_turnaround_seconds": round(
                avg_turnaround,
                2
            )
        })

    team_data.sort(
        key=lambda x: (
            x[sort_by]
            if x[sort_by] is not None
            else 0
        ),
        reverse=(
            sort_order == "desc"
        )
    )

    excel_rows = []

    for row in team_data:

        excel_rows.append([
            row["team_name"],
            row["member_count"],
            row["total_decisions"],
            row["approved_decisions"],
            row["rejected_decisions"],
            row["pending_decisions"],
            row["total_approvals"],
            row["approved_approvals"],
            row["rejected_approvals"],
            row["pending_approvals"],
            row["avg_turnaround_seconds"]
        ])

    summary = {

        "total_teams": len(team_data),

        "total_members": sum(
            r["member_count"]
            for r in team_data
        ),

        "total_decisions": sum(
            r["total_decisions"]
            for r in team_data
        ),

        "approved_decisions": sum(
            r["approved_decisions"]
            for r in team_data
        ),

        "rejected_decisions": sum(
            r["rejected_decisions"]
            for r in team_data
        ),

        "pending_decisions": sum(
            r["pending_decisions"]
            for r in team_data
        ),

        "total_approvals": sum(
            r["total_approvals"]
            for r in team_data
        )
    }

    excel = create_excel(
        "Team Report",
        {
            "team": team,
            "start_date": start_date,
            "end_date": end_date,
            "status": status,
            "category": category
        },
        summary,
        [
            "Team Name",
            "Member Count",
            "Total Decisions",
            "Approved Decisions",
            "Rejected Decisions",
            "Pending Decisions",
            "Total Approvals",
            "Approved Approvals",
            "Rejected Approvals",
            "Pending Approvals",
            "Avg Turnaround Seconds"
        ],
        excel_rows
    )

    return StreamingResponse(
        excel,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                "attachment; filename=team_report.xlsx"
        }
    )


# =========================================================
# AUDIT DATA HELPER
# =========================================================

def get_audit_data(
    db: Session,
    user: int | None = None,
    action: str | None = None,
    entity_type: str | None = None,
    entity_id: int | None = None,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
    sort_by: str = "created_at",
    sort_order: str = "desc"
):

    query = (
        db.query(AuditLog)
        .join(
            User,
            AuditLog.user_id == User.id
        )
    )

    if user:
        query = query.filter(
            AuditLog.user_id == user
        )

    if action:
        query = query.filter(
            AuditLog.action == action
        )

    if entity_type:
        query = query.filter(
            AuditLog.entity_type == entity_type
        )

    if entity_id:
        query = query.filter(
            AuditLog.entity_id == entity_id
        )

    if start_date:
        query = query.filter(
            AuditLog.created_at >= start_date
        )

    if end_date:
        query = query.filter(
            AuditLog.created_at <= end_date
        )

    sort_columns = {

        "id": AuditLog.id,

        "user_id": AuditLog.user_id,

        "action": AuditLog.action,

        "entity_type": AuditLog.entity_type,

        "entity_id": AuditLog.entity_id,

        "created_at": AuditLog.created_at
    }

    if sort_order == "asc":
        query = query.order_by(
            sort_columns[sort_by].asc()
        )
    else:
        query = query.order_by(
            sort_columns[sort_by].desc()
        )

    return query.all()


def build_audit_rows(logs):

    rows = []

    for log in logs:

        rows.append([
            log.user.full_name,
            log.action,
            log.entity_type,
            log.entity_id,
            log.description,
            log.created_at,
            log.ip_address
        ])

    return rows


# =========================================================
# AUDIT PDF EXPORT
# =========================================================

@router.get("/audit/export/pdf")
def export_audit_pdf(
    user: int | None = Query(default=None),
    action: str | None = Query(default=None),
    entity_type: str | None = Query(default=None),
    entity_id: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    require_admin(current_user)

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "user_id",
            "action",
            "entity_type",
            "entity_id",
            "created_at"
        }
    )

    if user is not None and user < 1:
        raise HTTPException(
            status_code=422,
            detail="user must be greater than or equal to 1"
        )

    if entity_id is not None and entity_id < 1:
        raise HTTPException(
            status_code=422,
            detail="entity_id must be greater than or equal to 1"
        )

    if action and action not in ALLOWED_AUDIT_ACTIONS:
        raise HTTPException(
            status_code=422,
            detail="Invalid action"
        )

    if (
        entity_type
        and entity_type not in ALLOWED_AUDIT_ENTITY_TYPES
    ):
        raise HTTPException(
            status_code=422,
            detail="Invalid entity_type"
        )

    validate_date_range(
        start_date,
        end_date
    )

    logs = get_audit_data(
        db,
        user,
        action,
        entity_type,
        entity_id,
        start_date,
        end_date,
        sort_by,
        sort_order
    )

    pdf = create_pdf(
        "Audit Report",
        {
            "user": user,
            "action": action,
            "entity_type": entity_type,
            "entity_id": entity_id,
            "start_date": start_date,
            "end_date": end_date
        },
        {
            "total": len(logs)
        },
        [
            "User",
            "Action",
            "Entity Type",
            "Entity ID",
            "Description",
            "Timestamp",
            "IP Address"
        ],
        build_audit_rows(logs)
    )

    return StreamingResponse(
        pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                "attachment; filename=audit_report.pdf"
        }
    )


# =========================================================
# AUDIT EXCEL EXPORT
# =========================================================

@router.get("/audit/export/excel")
def export_audit_excel(
    user: int | None = Query(default=None),
    action: str | None = Query(default=None),
    entity_type: str | None = Query(default=None),
    entity_id: int | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    require_admin(current_user)

    validate_sort(
        sort_by,
        sort_order,
        {
            "id",
            "user_id",
            "action",
            "entity_type",
            "entity_id",
            "created_at"
        }
    )

    if user is not None and user < 1:
        raise HTTPException(
            status_code=422,
            detail="user must be greater than or equal to 1"
        )

    if entity_id is not None and entity_id < 1:
        raise HTTPException(
            status_code=422,
            detail="entity_id must be greater than or equal to 1"
        )

    if action and action not in ALLOWED_AUDIT_ACTIONS:
        raise HTTPException(
            status_code=422,
            detail="Invalid action"
        )

    if (
        entity_type
        and entity_type not in ALLOWED_AUDIT_ENTITY_TYPES
    ):
        raise HTTPException(
            status_code=422,
            detail="Invalid entity_type"
        )

    validate_date_range(
        start_date,
        end_date
    )

    logs = get_audit_data(
        db,
        user,
        action,
        entity_type,
        entity_id,
        start_date,
        end_date,
        sort_by,
        sort_order
    )

    excel = create_excel(
        "Audit Report",
        {
            "user": user,
            "action": action,
            "entity_type": entity_type,
            "entity_id": entity_id,
            "start_date": start_date,
            "end_date": end_date
        },
        {
            "total": len(logs)
        },
        [
            "User",
            "Action",
            "Entity Type",
            "Entity ID",
            "Description",
            "Timestamp",
            "IP Address"
        ],
        build_audit_rows(logs)
    )

    return StreamingResponse(
        excel,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                "attachment; filename=audit_report.xlsx"
        }
    )